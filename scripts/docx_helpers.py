import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    """Set background color of a cell (hex string without #)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def add_callout_box(doc, title, text_items, bg_color="FDF2F4", border_color="C41230"):
    """Create a beautiful callout box with a colored left border and light background."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Custom left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B) if border_color=="C41230" else RGBColor(0x1E, 0x40, 0xAF)
    
    for item in text_items:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15
        
        if isinstance(item, tuple):
            prefix, content = item
            r_pre = p_item.add_run(prefix)
            r_pre.bold = True
            r_pre.font.name = 'Arial'
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            
            r_cnt = p_item.add_run(content)
            r_cnt.font.name = 'Arial'
            r_cnt.font.size = Pt(10)
            r_cnt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        else:
            r = p_item.add_run(item)
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions ready")
